"""Create editable article/SI/cover documents with native Word equations."""
from pathlib import Path
import re,json
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'submission';OUT.mkdir(exist_ok=True)

def runtext(p,text):
    # Strip markdown markers and preserve visible external-source destinations.
    text=re.sub(r'\[([^]]+)\]\(([^)]+)\)',lambda m:m.group(1)+' ('+m.group(2)+')',text)
    for part in re.split(r'(\*\*.*?\*\*)',text):
        r=p.add_run(part[2:-2] if part.startswith('**') else part.replace('`',''))
        r.bold=part.startswith('**')

def mr(text):
    r=OxmlElement('m:r');pr=OxmlElement('m:rPr');sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');pr.append(sty);r.append(pr)
    wp=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');wp.append(fonts)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');wp.append(sz);r.append(wp)
    t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=text;r.append(t);return r

def pack(tag,parts):
    e=OxmlElement('m:'+tag)
    for p in parts:e.append(mr(p) if isinstance(p,str) else p)
    return e

def sub(base,index):return pack('sSub',[pack('e',[base]),pack('sub',[index])])
def sup(base,index):return pack('sSup',[pack('e',[base]),pack('sup',[index])])
def fraction(top,bottom):return pack('f',[pack('num',top),pack('den',bottom)])
def summation(index,expr):
    n=OxmlElement('m:nary');pr=OxmlElement('m:naryPr')
    for tag,val in [('chr','∑'),('limLoc','subSup'),('supHide','1')]:
        x=OxmlElement('m:'+tag);x.set(qn('m:val'),val);pr.append(x)
    n.append(pr);n.append(pack('sub',[index]));n.append(pack('sup',[]));n.append(pack('e',expr));return n

def binomial(top,bottom):
    f=fraction([top],[bottom]);pr=OxmlElement('m:fPr');typ=OxmlElement('m:type');typ.set(qn('m:val'),'noBar');pr.append(typ);f.insert(0,pr)
    d=OxmlElement('m:d');dp=OxmlElement('m:dPr')
    for tag,val in [('begChr','('),('endChr',')')]:
        x=OxmlElement('m:'+tag);x.set(qn('m:val'),val);dp.append(x)
    d.append(dp);d.append(pack('e',[f]));return d

def equation_parts(number):
    st=lambda x:sub(x,'t')
    if number==1:return [[st('E'),' = 1{',summation('j',['1(',sub('Y','jt'),' > ',sub('q','jt'),')']),' ≥ 2},   ',st('p'),' = P(',st('E'),' = 1 | ',sub('ℱ','d'),').  (1)']]
    if number==2:return [['max ',summation('t ∈ d',[st('a'),'(v',st('p'),' − c)'])],[st('a'),' ∈ {0,1},  ',summation('t ∈ d',[st('a')]),' ≤ K.  (2)']]
    if number==3:return [[sub('T','s'),' = {t : ',sub('τ','s'),' ≤ t < ',sub('τ','s'),' + 6 hours, ',st('E'),' = 1},'],['C(A) = ',summation('s',['1{A ∩ ',sub('T','s'),' ≠ ∅}']),'.  (3)']]
    if number==4:return [[fraction([sub('U','process')],[sub('v','s')]),' = C(A) − ',sub('r','s'),'N(A),   ',fraction([sub('U','hour')],[sub('v','h')]),' = H(A) − ',sub('r','h'),'N(A).  (4)']]
    if number==5:return [[st('p'),' = ',summation('i < j',['P(',sub('B','i'),' = 1, ',sub('B','j'),' = 1)']),' − 2P(',sub('B','1'),' = ',sub('B','2'),' = ',sub('B','3'),' = 1)'],['= 1 − ',summation('i < j',[sub('C','ij'),'(',sub('u','i'),', ',sub('u','j'),')']),' + 2',sub('C','123'),'(',sub('u','1'),', ',sub('u','2'),', ',sub('u','3'),').  (5)']]
    if number==6:return [[st('b'),' = ',st('p'),'[1 − max(',sub('p','t−1'),', …, ',sub('p','t−6'),')].  (6)']]
    if number==7:return [['|p(u′) − p(u)| ≤ ',summation('j',['|',sub('u','j'),'′ − ',sub('u','j'),'|']),'.  (7)']]
    if number==8:return [[sup(st('p'),'−'),' ≥ g  and  #{j ≠ t within the day: ',sup(sub('p','j'),'+'),' ≥ ',sup(st('p'),'−'),'} < K.  (8)']]
    if number==9:return [['P(early coverage after retention) = 1 − ',fraction([binomial('N−m','B')],[binomial('N','B')]),'.  (9)']]
    if number==10:return [[sub('p','true'),' = wz + (1−w)(3',sup('z','2'),'−2',sup('z','3'),').  (10)']]
    return None

def mathline(doc,text):
    # Each equation is an editable OMML object. Unicode scripts are native math glyphs.
    # Long displays are split at meaningful clause boundaries, not rasterized.
    number=re.search(r'\((\d+)\)\s*$',text)
    structured=equation_parts(int(number.group(1))) if number else None
    if structured:
        for expr in structured:
            p=doc.add_paragraph();p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(4);p.paragraph_format.keep_together=True
            mp=pack('oMathPara',[pack('oMath',expr)]);p._p.append(mp)
        return
    parts=[text]
    if len(text)>118:
        if 'subject to' in text:
            x=text.index('subject to');parts=[text[:x].rstrip(', '),text[x:]]
        elif '=1−' in text:
            x=text.index('=1−');parts=[text[:x],text[x:]]
        elif '  ' in text:
            parts=text.split('  ',1)
    for part in parts:
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_together=True
        mp=OxmlElement('m:oMathPara');m=OxmlElement('m:oMath');r=OxmlElement('m:r')
        pr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');pr.append(fonts)
        sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');pr.append(sz);r.append(pr)
        t=OxmlElement('m:t');t.set(qn('xml:space'),'preserve');t.text=part;r.append(t);m.append(r);mp.append(m);p._p.append(mp)

def table(doc,lines):
    rows=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    rows=[r for r in rows if not all(re.fullmatch(r':?-+:?',c.replace(' ','')) for c in r)]
    n=len(rows[0]);t=doc.add_table(rows=1,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    # Text labels get more space than numeric fields.
    weights=[]
    for j in range(n):
        body=[r[j] for r in rows[1:] if len(r)>j]
        maxlen=max((len(x) for x in body),default=8)
        weights.append(min(30,max(8,maxlen if maxlen>15 else 10)))
    widths=[6.5*w/sum(weights) for w in weights]
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=False
            if rows[0] == ['zone','candidate','reference','difference','lower','upper'] and len(rows)<=4:
                p.paragraph_format.keep_with_next = i < len(rows)-1
            numeric=bool(re.fullmatch(r'[−–+\d.,/%()\[\] ]+',text))
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if numeric or i==0 else WD_ALIGN_PARAGRAPH.LEFT
            runtext(p,text)
            for r in p.runs:r.font.size=Pt(9);r.bold=i==0;r.font.color.rgb=RGBColor(0,0,0)
            tcPr=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            tcPr.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ('top','bottom','left','right'):
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
            tcPr.append(margins)
            if i==0:
                shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'F2F2F2');tcPr.append(shading)
        trPr=t.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if i==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

def build(source,dest):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for style in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        f=doc.styles[style].font;f.name='Times New Roman';f.color.rgb=RGBColor(0,0,0)
        for border in doc.styles[style]._element.xpath('.//w:pBdr'):
            border.getparent().remove(border)
    normal=doc.styles['Normal'];normal.font.size=Pt(11);normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(16);doc.styles['Title'].font.bold=True
    for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
    doc.styles['Caption'].font.size=Pt(10);doc.styles['Caption'].font.italic=False
    # Footer field, no decorative header or cover page.
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
    lines=source.read_text(encoding='utf-8').splitlines();i=0;mathcount=0;tables=0;figs=0
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('|') and line.endswith('|'):
            chunk=[line]
            while i<len(lines) and lines[i].strip().startswith('|'):chunk.append(lines[i].strip());i+=1
            table(doc,chunk);tables+=1;continue
        if re.match(r'^!\[',line):
            m=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line);path=(source.parent/m.group(2)).resolve()
            high=Path(str(path).replace('_preview.png','.png'))
            if high.exists():path=high
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(path),width=Inches(6.45));figs+=1;continue
        if re.search(r'\(\d+\)\s*$',line) and not line.startswith('#'):
            mathline(doc,line);mathcount+=1;continue
        if line.startswith('# '):
            p=doc.add_paragraph(style='Title');runtext(p,line[2:]);continue
        if line.startswith('## '):
            p=doc.add_heading('',level=1);runtext(p,line[3:]);continue
        if line.startswith('### '):
            p=doc.add_heading('',level=2);runtext(p,line[4:]);continue
        if line.startswith('#### '):
            p=doc.add_heading('',level=3);runtext(p,line[5:]);continue
        caption=bool(re.match(r'^(Table|Figure) S?\d+\.',line))
        p=doc.add_paragraph(style='Caption' if caption else 'Normal')
        if caption and line.startswith('Table'):p.paragraph_format.keep_with_next=True
        if caption:p.paragraph_format.keep_together=True
        if source.name=='asmbi_manuscript.md' and line.startswith('Table 1.'):
            p.paragraph_format.page_break_before=True
        runtext(p,line)
    doc.core_properties.author='Junjie Zhang';doc.core_properties.title=lines[0].lstrip('# ');doc.core_properties.subject='ASMBI research article revision'
    doc.save(dest)
    return {'source':str(source.relative_to(ROOT)),'output':str(dest.relative_to(ROOT)),'native_equations':mathcount,'tables':tables,'figures':figs}

if __name__=='__main__':
    import argparse
    a=argparse.ArgumentParser();a.add_argument('--only',choices=['main','supporting','cover','all'],default='all');args=a.parse_args()
    jobs=[('main','asmbi_manuscript.md','ASMBI_manuscript.docx'),('supporting','asmbi_supporting_information.md','ASMBI_supporting_information.docx'),('cover','asmbi_cover_letter.md','ASMBI_cover_letter.docx')]
    report=[]
    for kind,src,dst in jobs:
        if args.only not in (kind,'all'):continue
        report.append(build(ROOT/'manuscript'/src,OUT/dst))
    (ROOT/'audit'/f'document_build_{args.only}.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
    print(json.dumps(report,indent=2))
